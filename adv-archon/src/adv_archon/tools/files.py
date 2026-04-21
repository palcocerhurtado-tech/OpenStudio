"""File tools — read_file, list_dir, grep, find."""

from __future__ import annotations

import os
import subprocess
from pathlib import Path


def _expand(path: str) -> Path:
    return Path(path).expanduser().resolve()


def read_file(path: str, start_line: int = 0, end_line: int = 0) -> str:
    p = _expand(path)
    if not p.exists():
        return f"Error: archivo no encontrado: {path}"

    suffix = p.suffix.lower()

    # PDF
    if suffix == ".pdf":
        return _read_pdf(p)

    # DOCX
    if suffix == ".docx":
        return _read_docx(p)

    # XLSX
    if suffix == ".xlsx":
        return _read_xlsx(p)

    # PPTX
    if suffix == ".pptx":
        return _read_pptx(p)

    # HTML
    if suffix in (".html", ".htm"):
        return _read_html(p)

    # Plain text / code
    try:
        text = p.read_text(errors="replace")
    except Exception as e:
        return f"Error leyendo {path}: {e}"

    lines = text.splitlines()
    if start_line or end_line:
        s = max(0, start_line - 1)
        e = end_line if end_line else len(lines)
        lines = lines[s:e]
    return "\n".join(lines)


def list_dir(path: str = ".", depth: int = 1) -> str:
    p = _expand(path)
    if not p.exists():
        return f"Error: directorio no encontrado: {path}"
    if not p.is_dir():
        return f"Error: {path} no es un directorio"

    IGNORE = {".git", "node_modules", ".venv", "__pycache__", ".mypy_cache", "dist", "build"}
    lines: list[str] = []

    def _walk(current: Path, level: int) -> None:
        if level > depth:
            return
        try:
            entries = sorted(current.iterdir(), key=lambda x: (x.is_file(), x.name.lower()))
        except PermissionError:
            return
        for entry in entries:
            if entry.name in IGNORE or entry.name.startswith("."):
                continue
            indent = "  " * (level - 1)
            suffix = "/" if entry.is_dir() else ""
            lines.append(f"{indent}{entry.name}{suffix}")
            if entry.is_dir() and level < depth:
                _walk(entry, level + 1)

    _walk(p, 1)
    return "\n".join(lines) if lines else "(vacío)"


def grep(pattern: str, path: str = ".", flags: str = "") -> str:
    cmd = ["rg", "--line-number", "--max-count", "50"]
    if "i" in flags:
        cmd.append("-i")
    cmd += [pattern, _expand(path).__str__()]
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=15)
        return result.stdout.strip() or "Sin coincidencias."
    except FileNotFoundError:
        # fallback to grep
        cmd2 = ["grep", "-rn", "--include=*.*", "-m", "50"]
        if "i" in flags:
            cmd2.append("-i")
        cmd2 += [pattern, str(_expand(path))]
        result = subprocess.run(cmd2, capture_output=True, text=True, timeout=15)
        return result.stdout.strip() or "Sin coincidencias."
    except Exception as e:
        return f"Error: {e}"


def find_files(name_pattern: str, path: str = ".") -> str:
    base = _expand(path)
    import fnmatch
    matches: list[str] = []
    for root, dirs, files in os.walk(base):
        dirs[:] = [d for d in dirs if d not in {".git", "node_modules", ".venv", "__pycache__"}]
        for f in files:
            if fnmatch.fnmatch(f, name_pattern):
                full = Path(root) / f
                matches.append(str(full).replace(str(base), "."))
        if len(matches) > 50:
            break
    return "\n".join(matches) if matches else "Sin resultados."


# ---------------------------------------------------------------------------
# Format readers
# ---------------------------------------------------------------------------

def _read_pdf(p: Path) -> str:
    try:
        import pdfplumber  # type: ignore
        with pdfplumber.open(p) as pdf:
            pages = []
            for page in pdf.pages:
                text = page.extract_text() or ""
                if text:
                    pages.append(text)
        return "\n\n".join(pages) if pages else "[PDF sin texto extraíble]"
    except ImportError:
        pass
    try:
        from pypdf import PdfReader  # type: ignore
        reader = PdfReader(str(p))
        texts = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(t for t in texts if t)
    except Exception as e:
        return f"Error leyendo PDF: {e}"


def _read_docx(p: Path) -> str:
    try:
        from docx import Document  # type: ignore
        doc = Document(str(p))
        return "\n".join(para.text for para in doc.paragraphs if para.text.strip())
    except ImportError:
        return "python-docx no instalado. Ejecuta: uv pip install python-docx"
    except Exception as e:
        return f"Error leyendo DOCX: {e}"


def _read_xlsx(p: Path) -> str:
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        parts: list[str] = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                rows.append("\t".join(cells))
            if rows:
                parts.append(f"=== {sheet.title} ===\n" + "\n".join(rows))
        return "\n\n".join(parts)
    except ImportError:
        return "openpyxl no instalado. Ejecuta: uv pip install openpyxl"
    except Exception as e:
        return f"Error leyendo XLSX: {e}"


def _read_pptx(p: Path) -> str:
    try:
        from pptx import Presentation  # type: ignore
        prs = Presentation(str(p))
        slides: list[str] = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                slides.append(f"--- Diapositiva {i} ---\n" + "\n".join(texts))
        return "\n\n".join(slides)
    except ImportError:
        return "python-pptx no instalado. Ejecuta: uv pip install python-pptx"
    except Exception as e:
        return f"Error leyendo PPTX: {e}"


def _read_html(p: Path) -> str:
    try:
        import trafilatura  # type: ignore
        text = trafilatura.extract(p.read_text(errors="replace"))
        return text or p.read_text(errors="replace")
    except ImportError:
        return p.read_text(errors="replace")


# ---------------------------------------------------------------------------
# Tool definitions
# ---------------------------------------------------------------------------

TOOL_DEFINITIONS = [
    {
        "name": "read_file",
        "description": "Lee el contenido de un archivo local (texto, PDF, DOCX, XLSX, PPTX, HTML).",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Ruta al archivo"},
                "start_line": {"type": "integer", "default": 0},
                "end_line": {"type": "integer", "default": 0},
            },
            "required": ["path"],
        },
        "function": read_file,
    },
    {
        "name": "list_dir",
        "description": "Lista el contenido de un directorio local.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "default": "."},
                "depth": {"type": "integer", "default": 1},
            },
        },
        "function": list_dir,
    },
    {
        "name": "grep",
        "description": "Busca texto o patrón en archivos locales.",
        "parameters": {
            "type": "object",
            "properties": {
                "pattern": {"type": "string"},
                "path": {"type": "string", "default": "."},
                "flags": {"type": "string", "default": ""},
            },
            "required": ["pattern"],
        },
        "function": grep,
    },
    {
        "name": "find_files",
        "description": "Busca archivos por nombre o patrón glob.",
        "parameters": {
            "type": "object",
            "properties": {
                "name_pattern": {"type": "string"},
                "path": {"type": "string", "default": "."},
            },
            "required": ["name_pattern"],
        },
        "function": find_files,
    },
]
